#!/usr/bin/env python3
"""
Painel no navegador para a automação de notas fiscais.

Abre uma página em http://localhost:8765 onde você:
  - arrasta o XML da nota e ela é lançada na hora;
  - vê todas as parcelas com situação, totais e filtros;
  - marca uma parcela como paga com um clique;
  - acompanha o histórico do que foi processado.

A pasta "entrada" continua sendo vigiada: um XML salvo nela por e-mail ou
pelo Explorador também é lançado em poucos segundos.

Uso:  python painel.py     (Ctrl+C no terminal para encerrar)
"""

import datetime as dt
import json
import logging
import os
import sys
import threading
import time
import webbrowser
from collections import deque
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote

import automacao_nf as nf

PORTA = 8765
BASE = Path(__file__).resolve().parent
PAGINA = BASE / "painel.html"
TRAVA = threading.Lock()          # uma operação na planilha por vez
EVENTOS = deque(maxlen=200)       # histórico mostrado na página
LIMITE_UPLOAD = 5 * 1024 * 1024   # 5 MB por arquivo


class GuardaEventos(logging.Handler):
    """Guarda as mensagens da automação para a página mostrar."""

    def emit(self, record):
        texto = record.getMessage()
        nivel = ("erro" if record.levelno >= logging.ERROR
                 else "aviso" if record.levelno >= logging.WARNING or "ignorada" in texto else "ok")
        EVENTOS.append({"hora": dt.datetime.now().strftime("%d/%m %H:%M:%S"),
                        "nivel": nivel, "texto": texto})


def configurar_log():
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    fmt = logging.Formatter("%(asctime)s  %(message)s", datefmt="%d/%m %H:%M:%S")
    raiz = logging.getLogger()
    raiz.setLevel(logging.INFO)
    for h in (logging.FileHandler(nf.ARQUIVO_LOG, encoding="utf-8"), logging.StreamHandler(sys.stdout)):
        h.setFormatter(fmt)
        raiz.addHandler(h)
    nf.log.addHandler(GuardaEventos())


def carregar_historico():
    """Mostra as últimas linhas do log antigo ao abrir o painel."""
    if not nf.ARQUIVO_LOG.exists():
        return
    linhas = nf.ARQUIVO_LOG.read_text(encoding="utf-8", errors="replace").splitlines()[-40:]
    for l in linhas:
        hora, _, texto = l.partition("  ")
        if not texto or texto.startswith(("Vigiando", "Encerrado", "Painel")):
            continue
        nivel = "erro" if "não é" in texto or "ilegível" in texto else "aviso" if "aberta" in texto or "ignorada" in texto else "ok"
        EVENTOS.append({"hora": hora, "nivel": nivel, "texto": texto})


# ---------------------------------------------------------------- dados
def _data(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def situacao(venc, pago, hoje):
    if str(pago or "").strip().lower() == "sim":
        return "Paga"
    if venc is None:
        return "Verificar"
    if venc < hoje:
        return "Vencida"
    if (venc - hoje).days <= 7:
        return "Vence em 7 dias"
    return "A vencer"


def ler_planilha():
    if not nf.PLANILHA.exists():
        return []
    from openpyxl import load_workbook
    wb = load_workbook(nf.PLANILHA, read_only=True)
    ws = wb[nf.ABA]
    hoje = dt.date.today()
    linhas = []
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=12, values_only=True), start=2):
        if not r or not r[1]:
            continue
        venc = _data(r[6])
        linhas.append({
            "linha": i,
            "lancado": _data(r[0]).isoformat() if _data(r[0]) else None,
            "fornecedor": r[1], "cnpj": r[2], "nf": str(r[3] or ""),
            "emissao": _data(r[4]).isoformat() if _data(r[4]) else None,
            "parcela": str(r[5] or ""),
            "vencimento": venc.isoformat() if venc else None,
            "dias": (venc - hoje).days if venc else None,
            "valor": float(r[7] or 0),
            "situacao": situacao(venc, r[9], hoje),
            "chave": r[10], "arquivo": r[11],
        })
    wb.close()
    return linhas


def marcar_pago(linha: int, pago: bool):
    from openpyxl import load_workbook
    wb = load_workbook(nf.PLANILHA)
    ws = wb[nf.ABA]
    if linha < 2 or linha > ws.max_row or not ws.cell(row=linha, column=2).value:
        raise ValueError("parcela não encontrada")
    ws[f"{nf.C_PAGO}{linha}"] = "Sim" if pago else None
    wb.save(nf.PLANILHA)
    forn, num, parc = (ws.cell(row=linha, column=c).value for c in (2, 4, 6))
    nf.log.info("NF %s de %s, parcela %s: %s", num, forn, parc, "marcada como paga" if pago else "voltou para em aberto")


def contar_entrada():
    if not nf.PASTA_ENTRADA.exists():
        return 0
    return sum(1 for f in nf.PASTA_ENTRADA.iterdir() if f.is_file() and not f.name.startswith("."))


# ---------------------------------------------------------------- servidor
class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):  # silencia o log de cada requisição
        pass

    def _json(self, obj, status=200):
        corpo = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(corpo)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(corpo)

    def _erro(self, msg, status=400):
        self._json({"ok": False, "erro": msg}, status)

    def _corpo(self, limite=LIMITE_UPLOAD):
        n = int(self.headers.get("Content-Length") or 0)
        if n > limite:
            raise ValueError("arquivo maior que 5 MB")
        return self.rfile.read(n)

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            corpo = PAGINA.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(corpo)))
            self.end_headers()
            self.wfile.write(corpo)
        elif self.path == "/api/dados":
            try:
                with TRAVA:
                    linhas = ler_planilha()
                self._json({"ok": True, "hoje": dt.date.today().isoformat(), "linhas": linhas,
                            "eventos": list(EVENTOS)[::-1], "pendentes": contar_entrada(),
                            "planilha": nf.PLANILHA.name, "pasta": str(nf.PASTA_ENTRADA)})
            except PermissionError:
                self._erro("A planilha está aberta no Excel. Feche-a para o painel ler os dados.", 409)
            except Exception as e:
                self._erro(f"Não consegui ler a planilha: {e}", 500)
        else:
            self.send_error(404)

    def do_POST(self):
        # aceita só chamadas da própria página
        origem = self.headers.get("Origin")
        if origem and origem not in (f"http://localhost:{PORTA}", f"http://127.0.0.1:{PORTA}"):
            return self._erro("origem não permitida", 403)
        try:
            if self.path == "/api/enviar":
                nome = Path(unquote(self.headers.get("X-Arquivo", ""))).name
                if not nome:
                    return self._erro("arquivo sem nome")
                dados = self._corpo()
                nf.PASTA_ENTRADA.mkdir(exist_ok=True)
                destino = nf.PASTA_ENTRADA / nome
                if destino.exists():
                    destino = nf.PASTA_ENTRADA / f"{destino.stem}_{int(time.time())}{destino.suffix}"
                destino.write_bytes(dados)
                return self._json({"ok": True})
            if self.path == "/api/processar":
                antes = len(EVENTOS)
                with TRAVA:
                    n = nf.processar_pasta()
                novos = list(EVENTOS)[antes:]
                return self._json({"ok": True, "lancadas": n, "eventos": novos})
            if self.path == "/api/pagar":
                p = json.loads(self._corpo(10_000) or b"{}")
                with TRAVA:
                    marcar_pago(int(p["linha"]), bool(p["pago"]))
                return self._json({"ok": True})
            if self.path == "/api/exemplos":
                import gerar_notas_exemplo
                gerar_notas_exemplo.main()
                nf.log.info("Notas de exemplo criadas na pasta entrada.")
                return self._json({"ok": True})
            if self.path == "/api/abrir-planilha":
                if not nf.PLANILHA.exists():
                    return self._erro("A planilha ainda não existe. Envie uma nota primeiro.")
                if hasattr(os, "startfile"):
                    os.startfile(nf.PLANILHA)  # Windows
                    return self._json({"ok": True})
                return self._erro("Abrir a planilha por aqui só funciona no Windows.")
            if self.path == "/api/abrir-pasta":
                nf.PASTA_ENTRADA.mkdir(exist_ok=True)
                if hasattr(os, "startfile"):
                    os.startfile(nf.PASTA_ENTRADA)
                    return self._json({"ok": True})
                return self._erro("Abrir a pasta por aqui só funciona no Windows.")
            self.send_error(404)
        except PermissionError:
            self._erro("A planilha está aberta no Excel. Feche-a e tente de novo.", 409)
        except ValueError as e:
            self._erro(str(e))
        except Exception as e:
            nf.log.error("Erro inesperado: %s", e)
            self._erro(f"Erro inesperado: {e}", 500)


def vigiar():
    while True:
        time.sleep(nf.INTERVALO_VIGIA_SEG)
        if contar_entrada():
            try:
                with TRAVA:
                    nf.processar_pasta()
            except Exception as e:
                nf.log.error("Erro ao processar a pasta: %s", e)


def main():
    configurar_log()
    carregar_historico()
    for p in (nf.PASTA_ENTRADA, nf.PASTA_PROCESSADAS, nf.PASTA_ERROS):
        p.mkdir(exist_ok=True)
    try:
        servidor = ThreadingHTTPServer(("127.0.0.1", PORTA), Handler)
    except OSError:
        print(f"A porta {PORTA} já está em uso: o painel provavelmente já está aberto.")
        print(f"Acesse http://localhost:{PORTA} no navegador.")
        webbrowser.open(f"http://localhost:{PORTA}")
        return
    threading.Thread(target=vigiar, daemon=True).start()
    url = f"http://localhost:{PORTA}"
    print(f"\n  Painel aberto em {url}\n  Deixe esta janela aberta. Para encerrar, aperte Ctrl+C.\n")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        servidor.serve_forever()
    except KeyboardInterrupt:
        print("\nPainel encerrado.")


if __name__ == "__main__":
    main()
